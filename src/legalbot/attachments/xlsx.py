"""XLSX extraction via openpyxl."""

from __future__ import annotations

import io
from typing import Any


def extract_xlsx(data: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets: dict[str, list[list[Any]]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([c for c in row])
        sheets[name] = rows
    text_lines: list[str] = []
    for name, rows in sheets.items():
        text_lines.append(f"# {name}")
        for row in rows:
            text_lines.append("\t".join("" if v is None else str(v) for v in row))
    return {
        "method": "openpyxl",
        "text": "\n".join(text_lines),
        "data": {
            "sheet_count": len(sheets),
            "sheets": sheets,
        },
    }
