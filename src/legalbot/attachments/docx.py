"""DOCX text extraction via python-docx."""

from __future__ import annotations

import io
from typing import Any


def extract_docx(data: bytes) -> dict[str, Any]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    tables: list[list[list[str]]] = []
    for t in doc.tables:
        rows: list[list[str]] = []
        for row in t.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)

    text = "\n".join(paragraphs)
    for rows in tables:
        text += "\n\n" + "\n".join(" | ".join(c) for c in rows)

    return {
        "method": "python-docx",
        "text": text,
        "data": {
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "tables": tables,
        },
    }
