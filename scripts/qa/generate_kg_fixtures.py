"""Generate attachment fixtures for KG full-capability QA.

Writes DOCX (NDA terms), PNG (logo for analyze_image routing), and TXT noise file.
Run from repo root: python scripts/qa/generate_kg_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

# Allow `python scripts/qa/generate_kg_fixtures.py` from repo root.
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))

from fixture_content import (
    ACME_LOGO_PNG_BYTES,
    EXECUTED_NDA_DOCX_SECTIONS,
    PRIOR_NDA_DOCX_SECTIONS,
)

OUT = Path(__file__).resolve().parent / "fixtures"


def _write_docx(path: Path, sections: list[tuple[str, str]]) -> None:
    from docx import Document

    doc = Document()
    for text, kind in sections:
        if kind == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    _write_docx(OUT / "prior_nda.docx", PRIOR_NDA_DOCX_SECTIONS)
    _write_docx(OUT / "executed_nda.docx", EXECUTED_NDA_DOCX_SECTIONS)
    (OUT / "acme_logo.png").write_bytes(ACME_LOGO_PNG_BYTES)
    (OUT / "irrelevant_memo.txt").write_text(
        "Internal cafeteria menu update. No legal review required.",
        encoding="utf-8",
    )

    print(f"Wrote fixtures under {OUT}")


if __name__ == "__main__":
    main()
